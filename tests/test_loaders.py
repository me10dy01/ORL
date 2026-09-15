"""多格式文档加载器测试（md / markdown / txt / pdf / docx）。"""

from __future__ import annotations

import pytest

from app.exceptions import ValidationError
from app.rag import loaders


def _make_minimal_pdf(text: str) -> bytes:
    """生成含可提取文本的最小 PDF。"""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = ("BT /F1 24 Tf 72 720 Td (" + text + ") Tj ET").encode("latin-1")
    objs.append(
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    )
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_is_supported() -> None:
    for ext in (".md", ".markdown", ".txt", ".pdf", ".docx", ".MD", ".PDF"):
        assert loaders.is_supported(f"a{ext}")
    assert not loaders.is_supported("a.xlsx")
    assert not loaders.is_supported("noext")


def test_extract_txt(tmp_path) -> None:
    p = tmp_path / "a.txt"
    p.write_text("你好，星辰数码商城。", encoding="utf-8")
    assert loaders.extract_pages(p) == [(1, "你好，星辰数码商城。")]


def test_extract_md(tmp_path) -> None:
    p = tmp_path / "a.md"
    p.write_text("# 标题\n正文内容", encoding="utf-8")
    pages = loaders.extract_pages(p)
    assert pages[0][0] == 1
    assert "正文内容" in pages[0][1]


def test_extract_docx(tmp_path) -> None:
    import docx

    p = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("第一段内容")
    d.add_paragraph("第二段内容")
    d.save(str(p))
    pages = loaders.extract_pages(p)
    assert len(pages) == 1
    assert "第一段内容" in pages[0][1]
    assert "第二段内容" in pages[0][1]


def test_extract_pdf(tmp_path) -> None:
    p = tmp_path / "a.pdf"
    p.write_bytes(_make_minimal_pdf("Hello PDF"))
    assert loaders.extract_pages(p) == [(1, "Hello PDF")]


def test_unsupported_raises(tmp_path) -> None:
    p = tmp_path / "a.xlsx"
    p.write_bytes(b"x")
    with pytest.raises(ValidationError):
        loaders.extract_pages(p)
